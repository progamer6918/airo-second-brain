"""Read an Owner-designated XLSX; publish private, provenance-bound CSV and review files.

No network, Telegram, scheduler, or source-workbook writes. Requires openpyxl 3.1.5.
"""
import argparse
import csv
from datetime import date, datetime
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unicodedata
import warnings

import openpyxl
from openpyxl.utils.datetime import from_excel
from openpyxl.utils import get_column_letter

HEADERS = ('NAMA DEALER', 'NAMA', 'TANGGAL LAHIR', 'TANPA TAHUN')
PLACEHOLDER_YEARS = {1900, 2024}  # Retain raw value; never infer age from a placeholder.
FIELDS = ('NAMA DEALER', 'NAMA', 'BIRTH_DAY', 'BIRTH_MONTH', 'BIRTH_YEAR',
          'YEAR_KNOWN', 'SOURCE_FILE_ID', 'SOURCE_SHA256', 'SOURCE_SHEET',
          'SOURCE_ROW', 'SOURCE_CELLS', 'DUPLICATE_GROUP', 'REVIEW_FLAGS')


def key(text):
    return ' '.join(unicodedata.normalize('NFKC', text).casefold().split())


def blank(value):
    return value is None or not str(value).strip() or str(value).strip() == '-'


def parse(value, epoch):
    if blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            result = from_excel(value, epoch)
            return result.date() if isinstance(result, datetime) else None
        except (ValueError, OverflowError):
            return None
    text = str(value).strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d/%m', '%d-%m'):
        try:
            if fmt in ('%d/%m', '%d-%m'):
                return datetime.strptime(text + '/2024', fmt + '/%Y').date()
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def extract(source, file_id, modified_time):
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            matches = []
            for sheet in workbook:
                for rownum, row in enumerate(sheet.iter_rows(max_row=60), 1):
                    positions = {str(c.value).strip().upper(): index for index, c in enumerate(row) if c.value is not None}
                    if all(h in positions for h in HEADERS):
                        matches.append((sheet, rownum, positions))
            if len(matches) != 1:
                raise ValueError('Exactly one worksheet must match all birthday headers')
            sheet, header_row, positions = matches[0]
            stats = dict.fromkeys(('TOTAL_SOURCE_ROWS', 'VALID_BIRTHDAY_ROWS', 'MISSING_NAME_ROWS',
                'MISSING_DATE_ROWS', 'MISSING_BOTH_DATE_ROWS', 'MISSING_DEALER_ROWS',
                'INVALID_DATE_ROWS', 'CONFLICT_ROWS', 'UNKNOWN_BIRTH_YEAR_ROWS',
                'NORMALIZED_ONLY_ROWS', 'POSSIBLE_DUPLICATE_ROWS'), 0)
            records = []
            for rownum, row in enumerate(sheet.iter_rows(min_row=header_row + 1,
                    max_col=max(positions[h] for h in HEADERS) + 1), header_row + 1):
                cells = [row[positions[h]] for h in HEADERS]
                values = [c.value for c in cells]
                if all(blank(v) for v in values):
                    continue
                stats['TOTAL_SOURCE_ROWS'] += 1
                dealer, name, raw_birth, raw_normalized = values
                birth, normalized = parse(raw_birth, workbook.epoch), parse(raw_normalized, workbook.epoch)
                flags = []
                if blank(name):
                    stats['MISSING_NAME_ROWS'] += 1; flags.append('MISSING_NAME')
                if blank(dealer):
                    stats['MISSING_DEALER_ROWS'] += 1; flags.append('MISSING_DEALER')
                if blank(raw_birth):
                    stats['MISSING_DATE_ROWS'] += 1
                if birth is None and normalized is None:
                    stats['MISSING_BOTH_DATE_ROWS'] += 1; flags.append('NO_USABLE_DATE')
                if (not blank(raw_birth) and birth is None) or (not blank(raw_normalized) and normalized is None):
                    stats['INVALID_DATE_ROWS'] += 1; flags.append('INVALID_DATE')
                if birth and normalized and (birth.month, birth.day) != (normalized.month, normalized.day):
                    flags.append('DATE_COLUMN_CONFLICT')
                chosen = birth or normalized
                known = birth is not None and birth.year not in PLACEHOLDER_YEARS and birth.year > 1900
                if chosen and not known:
                    stats['UNKNOWN_BIRTH_YEAR_ROWS'] += 1
                if birth is None and normalized is not None:
                    stats['NORMALIZED_ONLY_ROWS'] += 1
                output = {'NAMA DEALER': '' if blank(dealer) else str(dealer),
                    'NAMA': '' if blank(name) else str(name),
                    'BIRTH_DAY': chosen.day if chosen else '', 'BIRTH_MONTH': chosen.month if chosen else '',
                    'BIRTH_YEAR': birth.year if known else '', 'YEAR_KNOWN': int(known),
                    'SOURCE_FILE_ID': file_id, 'SOURCE_SHA256': source_hash, 'SOURCE_SHEET': sheet.title,
                    'SOURCE_ROW': rownum, 'SOURCE_CELLS': ','.join(get_column_letter(positions[h] + 1) + str(rownum) for h in HEADERS),
                    'DUPLICATE_GROUP': '', 'REVIEW_FLAGS': ''}
                records.append({'output': output, 'flags': flags,
                    'original_cells': {h: v.isoformat() if isinstance(v, (date, datetime)) else v for h, v in zip(HEADERS, values)}})
            # Preserve every row; identities merely identify possible duplicates/conflicts for review.
            groups = {}
            for record in records:
                out = record['output']
                if out['NAMA'] and out['NAMA DEALER'] and out['BIRTH_DAY']:
                    groups.setdefault((key(out['NAMA DEALER']), key(out['NAMA'])), []).append(record)
            for members in groups.values():
                if len(members) < 2:
                    continue
                dates = {(x['output']['BIRTH_MONTH'], x['output']['BIRTH_DAY']) for x in members}
                years = {x['output']['BIRTH_YEAR'] for x in members if x['output']['YEAR_KNOWN']}
                conflict = len(dates) > 1 or len(years) > 1
                group = hashlib.sha256(','.join(str(x['output']['SOURCE_ROW']) for x in members).encode()).hexdigest()[:16]
                for record in members:
                    record['flags'].append('IDENTITY_DATE_CONFLICT' if conflict else 'POSSIBLE_DUPLICATE')
                    record['output']['DUPLICATE_GROUP'] = group
                    if not conflict:
                        stats['POSSIBLE_DUPLICATE_ROWS'] += 1
            valid, review = [], []
            for record in records:
                flags = record['flags']; out = record['output']
                out['REVIEW_FLAGS'] = ','.join(flags)
                if any('CONFLICT' in f for f in flags):
                    stats['CONFLICT_ROWS'] += 1
                excluded = any(f != 'POSSIBLE_DUPLICATE' for f in flags)
                (review if excluded else valid).append(record)
            stats['VALID_BIRTHDAY_ROWS'] = len(valid)
            manifest = {'schema': 1, 'SOURCE_FETCH': 'PASS', 'SOURCE_FILE_ID': file_id,
                'SOURCE_MODIFIED_TIME': modified_time, 'SOURCE_SHA256': source_hash,
                'BIRTHDAY_TAB': sheet.title, 'HEADER_ROW': header_row, **stats,
                'REVIEW_ROWS': len(review), 'DATA_AUTHORITY': 'OWNER_DESIGNATED_LATEST_SOURCE',
                'ACTIVATED': False, 'policy': {'placeholder_primary_birth_years': sorted(PLACEHOLDER_YEARS),
                    'normalized_column_never_supplies_birth_year': True, 'conflicts_quarantined': True,
                    'possible_duplicates_preserved': True, 'missing_primary_date_uses_valid_normalized_day_month_only': True},
                'parser_versions': {'openpyxl': openpyxl.__version__}}
            return manifest, valid, review, records
        finally:
            workbook.close()


def publish(source, output_root, file_id, modified_time):
    os.umask(0o077)
    source, output_root = Path(source), Path(output_root)
    if any(p.is_symlink() for p in (output_root, *output_root.parents)):
        raise ValueError('Private output path must not use symlinks')
    output_root.mkdir(parents=True, exist_ok=True, mode=0o700)
    if output_root.stat().st_mode & 0o077:
        raise ValueError('Output root must be private (0700)')
    manifest, valid, review, records = extract(source, file_id, modified_time)
    extractor_hash = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    folder = output_root / (manifest['SOURCE_SHA256'][:16] + '-' + extractor_hash[:12])
    stream = io.StringIO(newline=''); writer = csv.DictWriter(stream, fieldnames=FIELDS)
    writer.writeheader(); writer.writerows(record['output'] for record in valid)
    csv_bytes = stream.getvalue().encode('utf-8')
    manifest.update({'PRIVATE_CSV_CREATED': True, 'PRIVATE_CSV_SHA256': hashlib.sha256(csv_bytes).hexdigest(),
                     'EXTRACTOR_SHA256': extractor_hash})
    files = {'birthdays.csv': csv_bytes, 'manifest.json': json.dumps(manifest, indent=2).encode(),
             'review.private.json': json.dumps(review, ensure_ascii=False, indent=2).encode(),
             'provenance.private.json': json.dumps(records, ensure_ascii=False, indent=2).encode()}
    if folder.exists():
        if any(not (folder / name).is_file() or (folder / name).read_bytes() != data for name, data in files.items()):
            raise ValueError('Existing extraction differs; refusing overwrite')
    else:
        pending = Path(tempfile.mkdtemp(prefix='.pending-', dir=output_root))
        for name, data in files.items():
            with (pending / name).open('xb') as output:
                output.write(data); output.flush(); os.fsync(output.fileno())
        os.rename(pending, folder)
        fd = os.open(output_root, os.O_RDONLY | os.O_DIRECTORY)
        try:
            os.fsync(fd)
        finally:
            os.close(fd)
    return {**manifest, 'PRIVATE_DATASET_DIR': str(folder)}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    for flag in ('source', 'output-root', 'source-file-id', 'source-modified-time'):
        parser.add_argument('--' + flag, required=True)
    args = parser.parse_args()
    try:
        result = publish(args.source, args.output_root, args.source_file_id, args.source_modified_time)
    except Exception as exc:
        print(json.dumps({'DATA_MIGRATION': 'BLOCKED', 'error_type': type(exc).__name__}))
        raise SystemExit(1) from None
    print(json.dumps(result, indent=2))
